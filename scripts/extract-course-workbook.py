import json
import re
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def clean(value):
    if value is None:
        return None
    text = str(value).replace("\u00a0", " ").strip()
    return text or None


workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
result = {}

for sheet_name in ["Course index", "Course details", "Curriculum", "Source checks"]:
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_header(value) for value in next(rows)]
    output_rows = []
    for row_number, row in enumerate(rows, start=2):
        values = {header: clean(value) for header, value in zip(headers, row)}
        if any(value for value in values.values()):
            values["__row"] = row_number
            output_rows.append(values)
    result[sheet_name] = output_rows

print(json.dumps(result, ensure_ascii=False))
